import os
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import fitz  # PyMuPDF
import tkinter as tk
from tkinter import filedialog, messagebox
from openpyxl import Workbook, load_workbook
import six
import win32com.client as win32  # for handling .doc files using Microsoft Word

def convert_doc_to_docx(file_path):
    new_file_path = os.path.splitext(file_path)[0] + '.docx'
    word = win32.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(file_path)
        doc.SaveAs(new_file_path, FileFormat=16)  # 16 represents .docx format
        doc.Close()
    except Exception as e:
        raise RuntimeError(f"An error occurred while converting the file: {e}")
    finally:
        word.Quit()
    return new_file_path

def analyze_docx(file_path):
    doc = Document(file_path)
    content = []
    
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else 'N/A'
        alignment = para.alignment
        if alignment == WD_ALIGN_PARAGRAPH.LEFT:
            alignment_str = 'LEFT'
        elif alignment == WD_ALIGN_PARAGRAPH.CENTER:
            alignment_str = 'CENTER'
        elif alignment == WD_ALIGN_PARAGRAPH.RIGHT:
            alignment_str = 'RIGHT'
        elif alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
            alignment_str = 'JUSTIFY'
        else:
            alignment_str = 'N/A'
        
        font_names = []
        font_sizes = []
        bolds = []
        italics = []
        underlines = []
        
        for run in para.runs:
            font = run.font
            if font.name:
                font_names.append(font.name)
            if font.size:
                font_sizes.append(font.size.pt)
            bolds.append(font.bold if font.bold is not None else False)
            italics.append(font.italic if font.italic is not None else False)
            underlines.append(font.underline if font.underline is not None else False)
        
        font_name = max(set(font_names), key=font_names.count) if font_names else 'N/A'
        font_size = max(set(font_sizes), key=font_sizes.count) if font_sizes else 'N/A'
        bold = any(bolds)
        italic = any(italics)
        underline = any(underlines)
        
        indent_info = para.paragraph_format.left_indent
        left_indent = indent_info.pt if indent_info else 0
        
        space_before = para.paragraph_format.space_before
        space_before = space_before.pt if space_before else 0
        space_after = para.paragraph_format.space_after
        space_after = space_after.pt if space_after else 0
        
        line_spacing = para.paragraph_format.line_spacing
        if line_spacing:
            if hasattr(line_spacing, 'pt'):
                line_spacing = line_spacing.pt
            else:
                line_spacing = line_spacing
        else:
            line_spacing = 'N/A'
        
        is_numbered = 'Yes' if para.style.name.startswith('List') else 'No'
        
        content.append({
            'text': para.text,
            'style': style_name,
            'alignment': alignment_str,
            'font_name': font_name,
            'font_size': font_size,
            'bold': bold,
            'italic': italic,
            'underline': underline,
            'left_indent': left_indent,
            'space_before': space_before,
            'space_after': space_after,
            'line_spacing': line_spacing,
            'is_numbered': is_numbered,
        })
        
    df = pd.DataFrame(content)
    consolidated_df = df.groupby(['style', 'alignment', 'font_name', 'font_size', 'bold', 'italic', 'underline', 'left_indent', 'space_before', 'space_after', 'line_spacing', 'is_numbered']).agg({'text': ' '.join}).reset_index()
    return consolidated_df

def analyze_pdf(file_path):
    doc = fitz.open(file_path)
    content = []
    
    for page_num, page in enumerate(doc):
        page_dict = page.get_text("dict")
        for block in page_dict["blocks"]:
            if block["type"] == 0:
                for line in block["lines"]:
                    for span in line["spans"]:
                        text_content = span["text"].strip()
                        if not text_content:
                            continue
                        font_name = span["font"]
                        font_size = span["size"]
                        bold = 'Bold' in font_name or 'bold' in font_name
                        italic = 'Italic' in font_name or 'italic' in font_name
                        underline = False
                        
                        content.append({
                            'page_number': page_num + 1,
                            'text': text_content,
                            'font_name': font_name,
                            'font_size': font_size,
                            'bold': bold,
                            'italic': italic,
                            'underline': underline,
                            'alignment': 'N/A',
                            'style': 'N/A',
                            'left_indent': 'N/A',
                            'space_before': 'N/A',
                            'space_after': 'N/A',
                            'line_spacing': 'N/A',
                            'is_numbered': 'N/A',
                        })
        
    df = pd.DataFrame(content)
    consolidated_df = df.groupby(['font_name', 'font_size', 'bold', 'italic', 'underline']).agg({'text': ' '.join}).reset_index()
    return consolidated_df

def analyze_document(file_path):
    if file_path.endswith('.docx'):
        return analyze_docx(file_path)
    elif file_path.endswith('.doc'):
        docx_file = convert_doc_to_docx(file_path)
        return analyze_docx(docx_file)
    elif file_path.endswith('.pdf'):
        return analyze_pdf(file_path)
    else:
        raise ValueError("Unsupported file type. Please provide a .doc, .docx, or .pdf file.")

def analyze_documents(file_paths):
    all_data = []
    for file_path in file_paths:
        df = analyze_document(file_path)
        all_data.append((os.path.basename(file_path), df))
    
    # Let user choose where to save the output file
    output_file = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel files", "*.xlsx")],
        title="Choose where to save the analysis"
    )
    
    if not output_file:
        messagebox.showinfo("Operation Cancelled", "File save was cancelled.")
        return
    
    # Save analysis to the chosen Excel file
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        for file_name, data in all_data:
            data.to_excel(writer, sheet_name=os.path.splitext(file_name)[0][:31], index=False)
    
    messagebox.showinfo("Analysis Complete", f"Analysis saved to {output_file}")

def main():
    root = tk.Tk()
    root.withdraw()
    
    file_paths = filedialog.askopenfilenames(
        title="Select Word or PDF Documents",
        filetypes=[("Word Documents", "*.doc;*.docx"), ("PDF Files", "*.pdf")]
    )
    
    if not file_paths:
        messagebox.showinfo("No file selected", "No files were selected. Exiting.")
        return
    
    analyze_documents(file_paths)

if __name__ == "__main__":
    main()